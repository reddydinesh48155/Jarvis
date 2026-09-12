"""Text extraction and normalization for supported knowledge documents."""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass
from pathlib import PurePath


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".md", ".markdown"}


class UnsupportedDocumentType(ValueError):
    """Raised when an upload has a file type outside the Part 6 contract."""


@dataclass(frozen=True)
class ExtractedSection:
    text: str
    page_number: int | None = None
    section: str | None = None


def document_extension(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def clean_text(text: str) -> str:
    """Normalize control characters and whitespace without destroying paragraphs."""
    normalized = unicodedata.normalize("NFKC", text).replace("\x00", " ")
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"[ \t]+", " ", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _extract_pdf(data: bytes) -> list[ExtractedSection]:
    from pypdf import PdfReader

    sections: list[ExtractedSection] = []
    reader = PdfReader(io.BytesIO(data))
    for page_index, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        if text:
            sections.append(ExtractedSection(text=text, page_number=page_index, section=f"Page {page_index}"))
    return sections


def _extract_docx(data: bytes) -> list[ExtractedSection]:
    from docx import Document

    document = Document(io.BytesIO(data))
    sections: list[ExtractedSection] = []
    current_section = "Document"
    buffer: list[str] = []

    def flush() -> None:
        text = clean_text("\n".join(buffer))
        if text:
            sections.append(ExtractedSection(text=text, section=current_section))
        buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name.lower() if paragraph.style is not None else ""
        if style_name.startswith("heading"):
            flush()
            current_section = text[:255]
        else:
            buffer.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                buffer.append(" | ".join(cells))
    flush()
    return sections


def _extract_plain_text(data: bytes, markdown: bool = False) -> list[ExtractedSection]:
    text = clean_text(data.decode("utf-8-sig", errors="replace"))
    if markdown and text:
        sections: list[ExtractedSection] = []
        current_section = "Document"
        buffer: list[str] = []
        for line in text.splitlines():
            heading = re.match(r"^#{1,6}\s+(.+?)\s*#*$", line.strip())
            if heading:
                section_text = clean_text("\n".join(buffer))
                if section_text:
                    sections.append(ExtractedSection(text=section_text, section=current_section))
                buffer.clear()
                current_section = heading.group(1)[:255]
            else:
                buffer.append(line)
        section_text = clean_text("\n".join(buffer))
        if section_text:
            sections.append(ExtractedSection(text=section_text, section=current_section))
        return sections
    return [ExtractedSection(text=text, section="Document")] if text else []


def extract_document(filename: str, data: bytes) -> list[ExtractedSection]:
    """Extract searchable sections while preserving source locations."""
    extension = document_extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentType(
            f"Unsupported file type '{extension or 'unknown'}'. Supported types: PDF, TXT, DOCX, Markdown."
        )
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    return _extract_plain_text(data, markdown=extension in {".md", ".markdown"})
